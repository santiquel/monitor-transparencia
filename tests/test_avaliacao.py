"""Testes da avaliação: regras (rubrica) e golden set."""

from openpyxl import load_workbook

from transparencia.catalog.loader import carregar_catalogo, carregar_entidades
from transparencia.evaluation.golden import ENTIDADES_PADRAO, gerar_modelo
from transparencia.evaluation.regras import carregar_regras
from transparencia.settings import get_settings

# Os 19 critérios da PoC, por nível de complexidade (docs/00_plano_e_decisoes.md §3)
CRITERIOS_POC = {
    *("1.1", "1.2", "1.3", "2.9"),  # existência simples
    *("16.2", "16.3", "16.4"),  # acessibilidade
    *("4.1", "5.1", "7.1"),  # exige navegação
    *("4.5", "18.3"),  # dados estruturados
    *("12.1", "20.1", "19.5"),  # documentos
    *("4.6", "20.2"),  # atualização
    *("15.2", "18.1"),  # interpretação semântica
}


def _config():
    s = get_settings()
    return (
        carregar_regras(s.evaluation_rules_path),
        carregar_catalogo(s.criteria_catalog_path),
        carregar_entidades(s.entities_path),
    )


def test_regras_cobrem_exatamente_os_criterios_da_poc():
    regras, catalogo, _ = _config()
    assert set(regras.criterios) == CRITERIOS_POC
    assert set(regras.criterios) <= {c.codigo for c in catalogo.criterios}
    assert regras.codigos()[:4] == ["1.1", "1.2", "1.3", "2.9"]  # ordem natural


def test_modelo_do_golden_set(tmp_path):
    regras, catalogo, cadastro = _config()
    destino = tmp_path / "golden.xlsx"
    linhas = gerar_modelo(destino, regras, catalogo, cadastro, ENTIDADES_PADRAO)
    assert linhas == len(ENTIDADES_PADRAO) * len(CRITERIOS_POC)

    ws = load_workbook(destino)["rotulos"]
    cabecalho = [c.value for c in ws[1]]
    assert cabecalho[:2] == ["entidade", "criterio"] and "status" in cabecalho
    primeira = {k: c.value for k, c in zip(cabecalho, ws[2], strict=True)}
    assert primeira["entidade"] == "rio-de-janeiro" and primeira["criterio"] == "1.1"
    assert primeira["status"] is None  # é a pessoa quem preenche
