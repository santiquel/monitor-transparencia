"""Medição: compara os vereditos do sistema com o golden set (rótulos manuais).

Por que não basta a acurácia: no golden set a maioria dos rótulos é ATENDE.
Um "sistema" que respondesse sempre ATENDE já teria acurácia alta. Por isso o
relatório traz, lado a lado:

- **cobertura**: em quantos casos o sistema decidiu (NAO_AVALIADO é abstenção,
  não erro: o sistema admitiu que não viu o suficiente);
- **acurácia nos decididos** e a do **baseline** "sempre a classe majoritária";
- **kappa de Cohen**: concordância descontado o acaso (0 = acaso, 1 = perfeita);
- **precisão e recall por status**, com destaque para NAO_ATENDE;
- **acusação falsa** (rótulo ATENDE, sistema NAO_ATENDE): o erro mais grave
  para um órgão de controle, e **omissão** (o inverso);
- **calibração**: dos erros, quantos o próprio sistema marcou para revisão.

O cálculo (`medir`) é puro, sem banco: recebe pares (rótulo, veredito).
"""

import csv
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from transparencia.db.models import (
    CrawlRun,
    CriterionResult,
    Entity,
    EvaluationRun,
    StatusConformidade,
    StatusExecucao,
)

ROTULOS_VALIDOS = ("ATENDE", "ATENDE_PARCIALMENTE", "NAO_ATENDE")
NAO_SEI = "NAO_SEI"
ABSTENCAO = StatusConformidade.NAO_AVALIADO.value
ORDEM_SISTEMA = [s.value for s in StatusConformidade]  # colunas da matriz de confusão


@dataclass
class Rotulo:
    entidade: str
    criterio: str
    status: str
    url_evidencia: str | None = None
    observacao: str | None = None
    verificado_em: date | None = None


@dataclass
class Previsao:
    entidade: str
    criterio: str
    status: str
    confianca: float | None = None
    metodo: str | None = None
    justificativa: str | None = None
    revisao_humana: bool = False
    evidencia_url: str | None = None
    evidencia_trecho: str | None = None
    evaluation_run_id: int | None = None


@dataclass
class Par:
    rotulo: Rotulo
    previsao: Previsao | None  # None = o sistema não tem resultado para o critério

    @property
    def sistema(self) -> str:
        return self.previsao.status if self.previsao else ABSTENCAO

    @property
    def decidido(self) -> bool:
        return self.sistema != ABSTENCAO

    @property
    def acertou(self) -> bool:
        return self.decidido and self.sistema == self.rotulo.status


@dataclass
class Medicao:
    pares: list[Par]
    ignorados_nao_sei: int
    # derivados
    n: int = 0
    decididos: int = 0
    acertos: int = 0
    cobertura: float = 0.0
    acuracia: float | None = None
    baseline_classe: str | None = None
    baseline_acuracia: float | None = None
    kappa: float | None = None
    por_status: dict[str, dict] = field(default_factory=dict)
    matriz: dict[str, Counter] = field(default_factory=dict)
    acusacoes_falsas: list[Par] = field(default_factory=list)
    omissoes: list[Par] = field(default_factory=list)
    erros_sinalizados: int = 0
    confianca_media_acertos: float | None = None
    confianca_media_erros: float | None = None
    por_criterio: dict[str, dict] = field(default_factory=dict)

    @property
    def divergencias(self) -> list[Par]:
        return [p for p in self.pares if p.decidido and not p.acertou]

    @property
    def abstencoes(self) -> list[Par]:
        return [p for p in self.pares if not p.decidido]


# ------------------------------------------------------------ cálculo


def _media(valores: list[float]) -> float | None:
    return sum(valores) / len(valores) if valores else None


def kappa_cohen(pares: list[tuple[str, str]]) -> float | None:
    """Kappa de Cohen entre dois "avaliadores" (rótulo humano × sistema)."""
    n = len(pares)
    if n == 0:
        return None
    observada = sum(a == b for a, b in pares) / n
    humano = Counter(a for a, _ in pares)
    sistema = Counter(b for _, b in pares)
    esperada = sum(humano[k] * sistema[k] for k in humano) / (n * n)
    if esperada == 1:
        return None  # todos na mesma classe: kappa indefinido
    return (observada - esperada) / (1 - esperada)


def medir(rotulos: list[Rotulo], previsoes: list[Previsao]) -> Medicao:
    indice = {(p.entidade, p.criterio): p for p in previsoes}
    validos = [r for r in rotulos if r.status in ROTULOS_VALIDOS]
    m = Medicao(
        pares=[Par(r, indice.get((r.entidade, r.criterio))) for r in validos],
        ignorados_nao_sei=sum(r.status == NAO_SEI for r in rotulos),
    )
    decididos = [p for p in m.pares if p.decidido]
    m.n, m.decididos = len(m.pares), len(decididos)
    m.acertos = sum(p.acertou for p in decididos)
    m.cobertura = m.decididos / m.n if m.n else 0.0
    m.acuracia = m.acertos / m.decididos if m.decididos else None

    if decididos:
        classe, qtd = Counter(p.rotulo.status for p in decididos).most_common(1)[0]
        m.baseline_classe, m.baseline_acuracia = classe, qtd / len(decididos)
    m.kappa = kappa_cohen([(p.rotulo.status, p.sistema) for p in decididos])

    for status in ROTULOS_VALIDOS:
        previstos = sum(p.sistema == status for p in decididos)
        reais = sum(p.rotulo.status == status for p in decididos)
        certos = sum(p.sistema == status == p.rotulo.status for p in decididos)
        m.por_status[status] = {
            "rotulados": sum(p.rotulo.status == status for p in m.pares),
            "previstos": previstos,
            "precisao": certos / previstos if previstos else None,
            "recall": certos / reais if reais else None,
        }

    m.matriz = {s: Counter() for s in ROTULOS_VALIDOS}
    for p in m.pares:
        m.matriz[p.rotulo.status][p.sistema] += 1

    m.acusacoes_falsas = [
        p for p in decididos if p.rotulo.status == "ATENDE" and p.sistema == "NAO_ATENDE"
    ]
    m.omissoes = [p for p in decididos if p.rotulo.status == "NAO_ATENDE" and p.sistema == "ATENDE"]
    m.erros_sinalizados = sum(p.previsao.revisao_humana for p in m.divergencias)

    def confs(certo: bool) -> list[float]:
        return [
            p.previsao.confianca
            for p in decididos
            if p.acertou == certo and p.previsao.confianca is not None
        ]

    m.confianca_media_acertos = _media(confs(True))
    m.confianca_media_erros = _media(confs(False))

    grupos: dict[str, list[Par]] = defaultdict(list)
    for p in m.pares:
        grupos[p.rotulo.criterio].append(p)
    for codigo, ps in grupos.items():
        dec = [p for p in ps if p.decidido]
        m.por_criterio[codigo] = {
            "n": len(ps),
            "decididos": len(dec),
            "acertos": sum(p.acertou for p in dec),
        }
    return m


# ------------------------------------------------------------ entrada


def ler_rotulos(caminho: Path) -> list[Rotulo]:
    ws = load_workbook(caminho, read_only=True)["rotulos"]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = list(next(linhas))
    rotulos = []
    for valores in linhas:
        d = dict(zip(cabecalho, valores, strict=False))
        if not d.get("entidade"):
            continue
        verificado = d.get("verificado_em")
        if isinstance(verificado, datetime):
            verificado = verificado.date()
        rotulos.append(
            Rotulo(
                entidade=str(d["entidade"]).strip(),
                criterio=str(d["criterio"]).strip(),
                status=str(d.get("status") or "").strip().upper(),
                url_evidencia=d.get("url_evidencia"),
                observacao=d.get("observacao"),
                verificado_em=verificado if isinstance(verificado, date) else None,
            )
        )
    return rotulos


def escolher_avaliacoes(
    session: Session, slugs: list[str], ids: list[int] | None = None
) -> dict[str, EvaluationRun]:
    """Avaliação a medir por entidade: as indicadas em `ids`, ou a mais recente concluída."""
    consulta = (
        select(EvaluationRun, Entity.slug)
        .join(Entity, Entity.id == EvaluationRun.entity_id)
        .options(
            selectinload(EvaluationRun.results).selectinload(CriterionResult.criterion),
            selectinload(EvaluationRun.results).selectinload(CriterionResult.evidences),
        )
    )
    if ids:
        consulta = consulta.where(EvaluationRun.id.in_(ids))
    else:
        consulta = consulta.where(
            Entity.slug.in_(slugs),
            EvaluationRun.status.in_([StatusExecucao.SUCESSO, StatusExecucao.PARCIAL]),
        )
    escolhidas: dict[str, EvaluationRun] = {}
    for run, slug in session.execute(consulta.order_by(EvaluationRun.started_at.desc())):
        escolhidas.setdefault(slug, run)  # a primeira de cada entidade é a mais recente
    return escolhidas


def previsoes_da_avaliacao(slug: str, run: EvaluationRun) -> list[Previsao]:
    previsoes = []
    for r in run.results:
        ev = r.evidences[0] if r.evidences else None
        previsoes.append(
            Previsao(
                entidade=slug,
                criterio=r.criterion.codigo,
                status=r.status,
                confianca=float(r.confianca) if r.confianca is not None else None,
                metodo=r.metodo,
                justificativa=r.justificativa,
                revisao_humana=r.revisao_humana,
                evidencia_url=ev.url if ev else None,
                evidencia_trecho=ev.trecho if ev else None,
                evaluation_run_id=run.id,
            )
        )
    return previsoes


# ------------------------------------------------------------ relatório


def pct(x: float | None) -> str:
    return "—" if x is None else f"{100 * x:.0f}%"


def _num(x: float | None, casas: int = 2) -> str:
    return "—" if x is None else f"{x:.{casas}f}"


def _celula(texto: str | None, limite: int = 160) -> str:
    t = " ".join(str(texto or "").split()).replace("|", "\\|")
    return t if len(t) <= limite else t[: limite - 1] + "…"


def _ordem_criterio(codigo: str) -> tuple[int, ...]:
    return tuple(int(x) for x in codigo.split("."))


def gerar_relatorio(
    m: Medicao,
    nome: str,
    avaliacoes: dict[str, EvaluationRun],
    coletas: dict[int, CrawlRun],
    versao_regras_atual: str,
) -> str:
    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    linhas = [
        f"# Medição contra o golden set — {nome}",
        "",
        f"Gerado em {agora} por `transparencia golden medir`. Gabarito: "
        "`golden/golden_set.xlsx` (rótulos manuais, feitos antes de ver o sistema).",
        "",
        "## Avaliações medidas",
        "",
        "| Entidade | Avaliação | Coleta | Data da coleta | Versão das regras | Custo IA (US$) |",
        "|---|---|---|---|---|---|",
    ]
    for slug, run in sorted(avaliacoes.items()):
        coleta = coletas.get(run.crawl_run_id) if run.crawl_run_id else None
        data = coleta.started_at.strftime("%Y-%m-%d") if coleta and coleta.started_at else "—"
        linhas.append(
            f"| {slug} | {run.id} | {run.crawl_run_id or '—'} | {data} | {run.versao_regras} "
            f"| {_num(float(run.custo_llm_usd or 0), 5)} |"
        )
    versoes = {r.versao_regras for r in avaliacoes.values()}
    if versoes != {versao_regras_atual}:
        linhas += [
            "",
            f"> Atenção: as regras atuais são a versão `{versao_regras_atual}`, mas as "
            f"avaliações medidas usaram {sorted(versoes)}. Reavalie para medir as regras atuais.",
        ]
    modelos = Counter(
        p.previsao.metodo for p in m.pares if p.previsao and p.previsao.metodo == "REGRA_LLM"
    )
    ia = "sim" if modelos else "não (nenhum resultado com método REGRA_LLM)"

    linhas += [
        "",
        "## Resumo",
        "",
        "| Métrica | Valor | Como ler |",
        "|---|---|---|",
        f"| Casos rotulados | {m.n} | {m.ignorados_nao_sei} NAO_SEI fora da medição |",
        f"| Cobertura | {pct(m.cobertura)} ({m.decididos}/{m.n}) | casos em que o sistema "
        "decidiu; o resto é NAO_AVALIADO (abstenção, não erro) |",
        f"| Acurácia nos decididos | {pct(m.acuracia)} ({m.acertos}/{m.decididos}) | "
        "status idêntico ao rótulo |",
        f"| Baseline (sempre {m.baseline_classe}) | {pct(m.baseline_acuracia)} | o que um "
        "sistema sem inteligência alguma acertaria nos mesmos casos |",
        f"| Kappa de Cohen | {_num(m.kappa)} | concordância além do acaso: <0,4 fraca, "
        "0,4–0,6 moderada, 0,6–0,8 boa, >0,8 muito boa |",
        f"| Acusações falsas | {len(m.acusacoes_falsas)} | rótulo ATENDE, sistema NAO_ATENDE "
        "(o erro mais grave) |",
        f"| Omissões | {len(m.omissoes)} | rótulo NAO_ATENDE, sistema ATENDE |",
        f"| Erros já sinalizados p/ revisão | {m.erros_sinalizados}/{len(m.divergencias)} | "
        "divergências que o próprio sistema marcou como `revisao_humana` |",
        f"| Confiança média (acertos × erros) | {_num(m.confianca_media_acertos)} × "
        f"{_num(m.confianca_media_erros)} | confiança calibrada = maior nos acertos |",
        f"| IA usada | {ia} | |",
        "",
        "## Precisão e recall por status (nos decididos)",
        "",
        "| Status | Rotulados | Previstos | Precisão | Recall |",
        "|---|---|---|---|---|",
    ]
    for status, d in m.por_status.items():
        linhas.append(
            f"| {status} | {d['rotulados']} | {d['previstos']} | {pct(d['precisao'])} "
            f"| {pct(d['recall'])} |"
        )
    linhas += [
        "",
        "Precisão de NAO_ATENDE = das vezes em que o sistema acusou, quantas estavam certas.",
        "Recall de NAO_ATENDE = dos descumprimentos reais, quantos o sistema encontrou.",
        "",
        "## Matriz de confusão (linhas = rótulo humano, colunas = sistema)",
        "",
        "| Rótulo \\ Sistema | " + " | ".join(ORDEM_SISTEMA) + " |",
        "|---" * (len(ORDEM_SISTEMA) + 1) + "|",
    ]
    for status, cont in m.matriz.items():
        linhas.append(f"| **{status}** | " + " | ".join(str(cont[s]) for s in ORDEM_SISTEMA) + " |")

    linhas += [
        "",
        "## Por critério",
        "",
        "| Critério | Rotulados | Decididos | Acertos | Acurácia |",
        "|---|---|---|---|---|",
    ]
    for codigo in sorted(m.por_criterio, key=_ordem_criterio):
        d = m.por_criterio[codigo]
        acc = d["acertos"] / d["decididos"] if d["decididos"] else None
        linhas.append(f"| {codigo} | {d['n']} | {d['decididos']} | {d['acertos']} | {pct(acc)} |")

    def tabela_pares(titulo: str, pares: list[Par], explicacao: str) -> list[str]:
        saida = ["", f"## {titulo} ({len(pares)})", "", explicacao, ""]
        if not pares:
            return [*saida, "Nenhum."]
        saida += [
            "| Entidade | Crit. | Rótulo | Sistema | Conf. | Revisar | Justificativa do sistema "
            "| Observação do rotulador | Evidência humana | Evidência do sistema |",
            "|---|---|---|---|---|---|---|---|---|---|",
        ]
        for p in sorted(
            pares, key=lambda p: (_ordem_criterio(p.rotulo.criterio), p.rotulo.entidade)
        ):
            pv = p.previsao
            saida.append(
                f"| {p.rotulo.entidade} | {p.rotulo.criterio} | {p.rotulo.status} | {p.sistema} "
                f"| {_num(pv.confianca if pv else None)} | "
                f"{'sim' if pv and pv.revisao_humana else ''} "
                f"| {_celula(pv.justificativa if pv else 'sem resultado')} "
                f"| {_celula(p.rotulo.observacao)} | {_celula(p.rotulo.url_evidencia, 90)} "
                f"| {_celula(pv.evidencia_url if pv else None, 90)} |"
            )
        return saida

    linhas += tabela_pares(
        "Divergências",
        m.divergencias,
        "Casos em que o sistema decidiu diferente do rótulo. Para cada um, classificar a "
        "causa: erro de regra, limite da coleta, rubrica ambígua ou rótulo a rever.",
    )
    linhas += tabela_pares(
        "Abstenções (NAO_AVALIADO)",
        m.abstencoes,
        "Casos em que o sistema não decidiu. Não contam como erro, mas cada um é "
        "trabalho manual: reduzir abstenções sem criar erros é o objetivo.",
    )
    return "\n".join(linhas) + "\n"


def gravar_csv(m: Medicao, destino: Path) -> None:
    """Uma linha por caso: para análise em planilha/pandas."""
    colunas = [
        "entidade",
        "criterio",
        "rotulo",
        "sistema",
        "decidido",
        "acertou",
        "confianca",
        "metodo",
        "revisao_humana",
        "evaluation_run_id",
        "justificativa",
        "observacao_rotulador",
        "url_evidencia_humana",
        "url_evidencia_sistema",
        "trecho_sistema",
    ]
    with destino.open("w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: Excel lê acentos
        w = csv.writer(f)
        w.writerow(colunas)
        for p in m.pares:
            pv = p.previsao
            w.writerow(
                [
                    p.rotulo.entidade,
                    p.rotulo.criterio,
                    p.rotulo.status,
                    p.sistema,
                    p.decidido,
                    p.acertou,
                    pv.confianca if pv else None,
                    pv.metodo if pv else None,
                    pv.revisao_humana if pv else None,
                    pv.evaluation_run_id if pv else None,
                    pv.justificativa if pv else None,
                    p.rotulo.observacao,
                    p.rotulo.url_evidencia,
                    pv.evidencia_url if pv else None,
                    pv.evidencia_trecho if pv else None,
                ]
            )
