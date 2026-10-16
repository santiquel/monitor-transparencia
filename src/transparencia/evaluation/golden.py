"""Golden set: rótulos feitos à mão, usados como gabarito para medir os avaliadores.

Fluxo:
1. `transparencia golden modelo` gera `golden/golden_set.xlsx`, uma linha por
   (entidade, critério), com a rubrica ao lado.
2. A pessoa abre cada site/portal, decide o status pela rubrica e anota a URL
   da evidência. Isso deve ser feito ANTES de olhar a saída do sistema, para o
   gabarito não ser influenciado por ela.
3. Na etapa de medição, o sistema compara os vereditos com estes rótulos.

A planilha vai para o Git: o gabarito faz parte da metodologia e precisa ser
reprodutível.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from transparencia.catalog.schema import CadastroEntidades, Catalogo
from transparencia.evaluation.regras import RegrasAvaliacao

# Os 5 municípios do escopo principal (Niterói e Maricá são extensão opcional).
ENTIDADES_PADRAO = [
    "rio-de-janeiro",
    "duque-de-caxias",
    "japeri",
    "itaborai",
    "cachoeiras-de-macacu",
]

STATUS_ROTULO = ["ATENDE", "ATENDE_PARCIALMENTE", "NAO_ATENDE", "NAO_SEI"]

COLUNAS = [
    # (cabeçalho, largura, preenchido pela pessoa?)
    ("entidade", 20, False),
    ("criterio", 9, False),
    ("descricao", 45, False),
    ("rubrica", 60, False),
    ("onde_olhar", 35, False),
    ("site", 30, False),
    ("portal", 30, False),
    ("status", 22, True),
    ("url_evidencia", 40, True),
    ("observacao", 40, True),
    ("verificado_em", 14, True),
]

INSTRUCOES = """\
COMO ROTULAR

1. Rotular ANTES de ver o resultado do sistema. Vê-lo antes deixaria o gabarito
   influenciado e tiraria o valor da medição de acurácia.
2. Uma linha = um critério numa entidade. Abrir o site/portal (colunas F e G),
   procurar onde a coluna "onde_olhar" indica e decidir pela coluna "rubrica".
3. status: escolher na lista.
   - ATENDE / ATENDE_PARCIALMENTE / NAO_ATENDE: conforme a rubrica.
     ATENDE_PARCIALMENTE só existe quando a rubrica o define.
   - NAO_SEI: não foi possível decidir (site fora do ar, informação ambígua).
     Essas linhas ficam fora da medição; o motivo vai em "observacao".
4. url_evidencia: a página onde a informação foi encontrada (ou procurada, sem sucesso).
5. observacao: o que foi visto, em uma frase. Ex.: "tem lista de licitações, mas sem
   o vencedor". É o que permite entender uma divergência com o sistema depois.
6. verificado_em: a data da verificação (o portal muda com o tempo; a
   comparação justa é com a coleta mais próxima dessa data).

Dica: para "não atende", procurar com cuidado por cerca de 2 minutos (menu,
busca do site, mapa do site). Sem encontrar nesse tempo, é NAO_ATENDE: o
cidadão também não acharia.
"""


def gerar_modelo(
    destino: Path,
    regras: RegrasAvaliacao,
    catalogo: Catalogo,
    cadastro: CadastroEntidades,
    slugs: list[str],
) -> int:
    """Cria a planilha de rotulagem. Devolve o número de linhas a rotular."""
    entidades = {e.slug: e for e in cadastro.entidades}
    criterios = {c.codigo: c for c in catalogo.criterios}

    wb = Workbook()
    ws = wb.active
    ws.title = "rotulos"
    ws.append([nome for nome, _, _ in COLUNAS])

    cabecalho = Font(bold=True, color="FFFFFF")
    for i, (_, largura, da_pessoa) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i)
        celula.font = cabecalho
        cor = "2E7D32" if da_pessoa else "37474F"  # verde = preenchido por quem rotula
        celula.fill = PatternFill("solid", fgColor=cor)
        ws.column_dimensions[celula.column_letter].width = largura

    for slug in slugs:
        e = entidades[slug]
        for codigo in regras.codigos():
            regra = regras.criterios[codigo]
            rubrica = "\n".join(f"{status}: {texto}" for status, texto in regra.rubrica.items())
            ws.append(
                [
                    slug,
                    codigo,
                    criterios[codigo].descricao,
                    rubrica,
                    regra.onde_olhar,
                    str(e.url_site),
                    str(e.url_portal) if e.url_portal else "",
                    None,
                    None,
                    None,
                    None,
                ]
            )

    linhas = ws.max_row - 1
    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")

    lista = DataValidation(type="list", formula1=f'"{",".join(STATUS_ROTULO)}"', allow_blank=True)
    lista.error = "Escolha um status da lista."
    ws.add_data_validation(lista)
    lista.add(f"H2:H{ws.max_row}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    instr = wb.create_sheet("instrucoes")
    instr.column_dimensions["A"].width = 100
    for texto in INSTRUCOES.splitlines():
        instr.append([texto])
    instr.append([])
    instr.append([f"Régua: config/evaluation_rules.yaml, versão {regras.versao}"])

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return linhas
